"""Read the derivation-likely sheets in SU625.xlsx."""

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

fpath = r"C:\Users\drwho\OneDrive\Documents\SU625\SU625.xlsx"

# Sheets most likely to contain the derivation
sheets_to_read = [
    "Physical Meaning",
    "Spacetime Density Model",
    "Correlation Analysis",
    "Variable Light Speed",
    "SU Core Formula Test",
    "SU Factored Summary",
]

# Load with both formulas (data_only=False) and computed values (data_only=True)
wb_formulas = openpyxl.load_workbook(fpath, data_only=False)
wb_values = openpyxl.load_workbook(fpath, data_only=True)

for sheet_name in sheets_to_read:
    if sheet_name not in wb_formulas.sheetnames:
        print(f"--- {sheet_name}: NOT FOUND ---\n")
        continue

    print("=" * 78)
    print(f"SHEET: {sheet_name}")
    print("=" * 78)
    ws_f = wb_formulas[sheet_name]
    ws_v = wb_values[sheet_name]
    rows = min(ws_f.max_row, 80)
    cols = min(ws_f.max_column, 14)

    for row_idx in range(1, rows + 1):
        row_vals = []
        any_content = False
        for col_idx in range(1, cols + 1):
            cf = ws_f.cell(row=row_idx, column=col_idx)
            cv = ws_v.cell(row=row_idx, column=col_idx)
            v = cv.value
            f = cf.value
            if v is None and f is None:
                cell_str = ""
            elif isinstance(f, str) and f.startswith("="):
                # formula cell — show both formula and value
                val_str = f"{v:.4g}" if isinstance(v, float) else str(v)
                cell_str = f"{f}={val_str}"
            elif isinstance(v, float):
                cell_str = f"{v:.6g}"
            else:
                cell_str = str(v if v is not None else f)[:50]
            row_vals.append(cell_str)
            if cell_str != "":
                any_content = True
        if any_content:
            print(f"R{row_idx}: " + " | ".join(row_vals))
    print()
